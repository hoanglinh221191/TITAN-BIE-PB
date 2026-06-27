from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


REPO = Path(__file__).resolve().parents[1]
PACKAGE = REPO / "docs" / "fagbi_guarded_p4_package" / "submission_preview_synced"
INPUT = PACKAGE / "titan_fagbi_language_edited_clean.docx"
OUTPUT = PACKAGE / "titan_fagbi_language_edited_clean_repro_ablation.docx"
QA_CSV = PACKAGE / "fagbi_scope_z4w_docx_qa.csv"


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "D1D5DB")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_paragraph_after_table(table, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    p = Paragraph(new_p, table._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_table_after(paragraph, rows):
    doc = paragraph._parent
    table = doc.add_table(rows=len(rows), cols=len(rows[0]), width=Inches(7.3))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_text(cell, str(text), bold=(r == 0))
            if r == 0:
                shade_cell(cell, "E5E7EB")
    set_table_borders(table)
    paragraph._p.addnext(table._tbl)
    return table


def find_paragraph(doc: Document, needle: str):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise RuntimeError(f"Unable to find paragraph containing: {needle}")


def add_methods_table(doc: Document):
    anchor = find_paragraph(doc, "All reported TITAN benchmarks were executed")
    p = insert_paragraph_after(
        anchor,
        "Table M1. Reproducibility parameters for the manuscript-facing TITAN benchmarks.",
    )
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(9)

    rows = [
        ("Category", "Parameter", "Value used in the reported rows", "Evidence and notes"),
        ("Electrostatics", "Dielectric constants", "solute epsilon = 1; solvent epsilon = 80", "arg.in: -pbe-int-diel=1, -pbe-ext-diel=80"),
        ("Electrostatics", "Ionic strength and temperature", "0.15 M salt, 298.15 K", "arg.in: -pbe-salt=0.15, -pbe-temp=298.15"),
        ("Electrostatics", "Screening parameter", "kappa = 0.12746 A^-1 in TABI profile at 0.15 M", "Reported as kappa_inv_A in pbe_tabi_profile.csv"),
        ("Molecular preparation", "Charge/radius model", "AMBER/PDB2PQR charges and radii; TITAN force-field tag AMBER99SB", "PQR provenance plus -ff=AMBER99SB runtime option"),
        ("Molecular preparation", "PDB2PQR preparation", "PDB2PQR 3.7.1, AMBER, pH 7.0, chain retained, PROPKA titration-state method", "Fixture manifests record command and SHA256 checksums"),
        ("Molecular preparation", "Radius floor", "0.6000 A post-processing floor for zero-radius PDB2PQR atoms", "Recorded as postprocess-radius-floor=0.6000A"),
        ("Surface", "Boundary surface", "van der Waals molecular surface", "arg.in: -pbe-surface=vdw"),
        ("Surface", "Panel generation", "TABI sphere-patch triangular surface, surface-points request = 8", "Actual panel counts reported per protein in the system table"),
        ("Surface", "Collocation/discretization", "centroid panel collocation with analytic-singular near integration", "profile: discretization=centroid, near_integration=analytic_singular"),
        ("Linear solver", "Operator", "JCP2013 coupled two-row TABI boundary-integral system", "profile: bie_system=linear_pb_2x2"),
        ("Linear solver", "Direct TABI reference meaning", "direct summation of the dense panel-panel boundary-integral operator", "No grid or FMM acceleration in direct TABI rows"),
        ("Linear solver", "GMRES controls", "relative tolerance 1e-6, dphi 1e-6, restart 128, max iterations 720", "arg.in: -pbe-rms, -pbe-dphi, -pbe-tabi-gmres-restart, -pbe-max-iter"),
        ("Preconditioning", "Guarded p4 preconditioner", "sparse near-field ILU(0), neighbor cap 64, drop tolerance 0", "arg.in: -pbe-tabi-precond=fmm-nearfield-ilu0"),
        ("Verification", "Direct-operator audit", "512 sampled targets; audit RMS and sampled direct true residual gate <= 8e-3", "arg.in: -pbe-tabi-matvec-audit=sampled-detail"),
        ("Guarded p4", "Far-field policy", "Taylor order 4, theta 0.6, static tail tolerance 0.002, strict geometry guard theta 0.5", "Z4N.4/Z4O/Z4V guarded p4 policy"),
        ("Implementation", "Build and precision", "native MinGW build, AVX2 CPU path, double precision", "Build command used in gates: mingw32-make PROFILE=native TARGET=titan.exe"),
        ("Implementation", "Threads", "4 TITAN worker threads", "arg.in: -nt=4"),
        ("Benchmark environment", "Hardware", "Intel Core Ultra 7 255H workstation, 30.9 GB RAM, Windows", "Reported benchmark environment"),
        ("Benchmark environment", "Thread pinning and power mode", "OS scheduler default; no explicit thread pinning or power-mode lock recorded", "Limitation for external timing reproduction"),
        ("Benchmark environment", "Timer and RSS", "wall-clock process time; peak RSS from process working-set telemetry", "PowerShell runner records wall_seconds and peak_working_set_mb"),
    ]
    table = insert_table_after(p, rows)
    widths = [1.15, 1.55, 2.15, 2.35]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def insert_after_text(doc: Document, anchor_text: str, blocks):
    anchor = find_paragraph(doc, anchor_text)
    current = anchor
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            current = insert_paragraph_after_table(current, block[1], style="Heading 2") if hasattr(current, "_tbl") else insert_paragraph_after(current, block[1], style="Heading 2")
        elif kind == "paragraph":
            current = insert_paragraph_after_table(current, block[1]) if hasattr(current, "_tbl") else insert_paragraph_after(current, block[1])
            current.paragraph_format.space_after = Pt(6)
        elif kind == "caption":
            current = insert_paragraph_after_table(current, block[1]) if hasattr(current, "_tbl") else insert_paragraph_after(current, block[1])
            current.runs[0].bold = True
            current.runs[0].font.size = Pt(9)
            current.paragraph_format.space_before = Pt(8)
            current.paragraph_format.space_after = Pt(3)
        elif kind == "table":
            table = insert_table_after(current, block[1])
            widths = block[2]
            for row in table.rows:
                for idx, width in enumerate(widths):
                    row.cells[idx].width = Inches(width)
            current = table
        else:
            raise ValueError(kind)


def add_results_sensitivity(doc: Document):
    blocks = [
        ("heading", "3.5. Ablation and Sensitivity"),
        (
            "paragraph",
            "The manuscript-facing ladder already contains one important ablation: direct dense TABI, degree-one FMM with diagonal or local block preconditioning, degree-one FMM with sparse near-field ILU(0), and guarded fourth-order p4 with the same ILU(0) graph. This sequence separates the cost of dense panel-panel summation, the Krylov stalling that remains after a low-order FMM matvec acceleration, and the stabilization provided by sparse near-field ILU(0).",
        ),
        (
            "paragraph",
            "The ILU(0) setup overhead is small relative to the reduction in Krylov work. The near-source graph is extracted and factored once before GMRES, while the successful full-protein rows complete in 128 iterations. In the 1A63 guarded p4 row, the profile records 5.15 s for ILU(0) setup and 6.37 s for preconditioner application over the solve, compared with 151.9 s spent in p4 matvec application and 312.6 s total wall time. Thus the graph-based factorization does not erase the benefit of reducing the effective Krylov stiffness.",
        ),
        (
            "paragraph",
            "A second ablation is the operator guard. In the six-protein Z4W sensitivity matrix, turning off the strict geometry guard left five systems within the sampled direct-operator gate but caused full 1A63 to fail verification: internal GMRES converged in 128 iterations, yet the sampled audit RMS was 5.18e-2 and the sampled direct true residual was 5.20e-2. Restoring the strict geometric MAC guard at theta = 0.5 reduced the 1A63 audit to 2.53e-3. This is why the final p4 path is described as guarded rather than merely high-order.",
        ),
        (
            "paragraph",
            "The same Z4W matrix also shows that the accepted p4 policy is not a fragile theta choice. Opening-theta values 0.4, 0.5, 0.6, and 0.7 all passed the sampled direct-operator gate on all six proteins; the median wall time decreased modestly from 142.9 s at theta = 0.4 to 126.2 s at theta = 0.7. We retain theta = 0.6 as the manuscript policy because it was the locked validation setting from the repeatability gates and because the strict geometry guard, rather than the opening angle alone, is the accuracy-critical control.",
        ),
        ("caption", "Table 6. Current ablation and sensitivity coverage in the manuscript evidence package."),
        (
            "table",
            [
                ("Sensitivity axis", "Current evidence", "Manuscript-facing interpretation", "Remaining extension"),
                ("Direct vs FMM vs ILU(0)", "Z4V six-protein performance ladder", "Direct dense TABI is runtime-limited; low-order FMM accelerates matvecs but can fail direct-operator audit; sparse ILU(0) stabilizes GMRES.", "None for the main claim."),
                ("Guarded p4 vs degree-one FMM", "Z4N.4/Z4V/Z4W sampled direct audits", "Guarded p4 restores direct-operator agreement on all six full proteins.", "Independent full direct TABI energy remains future work for large systems."),
                ("Opening theta", "Z4W theta 0.4/0.5/0.6/0.7 on six proteins", "All four opening angles passed; median wall time ranged from 142.9 s to 126.2 s.", "A larger protein panel could revisit the speed/guard trade-off."),
                ("Geometry guard", "Z4W guard on/off on six proteins", "Removing the guard failed only 1A63, with audit RMS 5.18e-2; guard-on recovered 2.53e-3.", "The guard is retained as an accuracy contract, not only a safety margin."),
                ("Tail indicator", "Z4N.1-Z4N.3 plus Z4W fail-closed tail-off row", "The accepted gmres-p4 path requires static operator-tail refinement at tolerance 0.002.", "Tail-off is intentionally rejected by the guarded p4 validator."),
                ("Expansion order p", "Z4W p2/p3/p4/p5 rows", "Only p4 is valid for guarded gmres-p4; p2/p5 are unsupported and p3 is fail-closed by the policy validator.", "Separate non-guarded single-matvec order studies can be kept in supplement if needed."),
                ("ILU neighbor cap", "Z4W caps 32/64/128", "The manuscript guarded p4 validator accepts ILU neighbor cap 64; caps 32/128 are fail-closed for this policy.", "A future preconditioner-only study can relax validation deliberately."),
                ("Setup / solve / audit split", "profile timing fields and Z4W matrix", "Setup, matvec, preconditioner, and sampled-audit costs are recorded separately.", "Use the Z4W CSVs to build an extended supplement timing table."),
                ("Memory vs panel count", "Z4V/Z4W RSS fields", "Peak RSS stays below 512 MB for the six-protein panel.", "A larger public panel is the natural next memory-scaling test."),
                ("Residual curves", "Barnase curve plus Z4W residual traces", "The central convergence profile is shown for Barnase; Z4W exports traces for all six proteins.", "A supplement figure can show all six traces without overcrowding the main text."),
            ],
            [1.45, 1.7, 2.4, 1.55],
        ),
        (
            "paragraph",
            "This ablation table deliberately separates scientific failures from fail-closed policy checks. The only measured accuracy gap in Z4W is the expected 1A63 geometry-off row. The p = 3, tail-off, and ILU-cap perturbations are rejected before solving because the guarded p4 CLI requires the validated policy ingredients. We therefore report them as reproducibility-boundary checks rather than as numerical counterexamples.",
        ),
    ]
    insert_after_text(
        doc,
        "Within these limits, the results support a focused methodological conclusion:",
        blocks,
    )


def add_direct_tabi_sentence(doc: Document):
    p = find_paragraph(doc, "Direct dense TABI is useful as a conceptual baseline")
    if "Direct TABI denotes direct summation" not in p.text:
        p.add_run(
            " Here, Direct TABI denotes direct summation of the dense JCP panel-panel boundary-integral operator, not a grid PB solve and not a treecode/FMM approximation."
        )


def save_markdown_note():
    md = PACKAGE / "titan_fagbi_language_edited_clean_repro_ablation_notes.md"
    md.write_text(
        "# Reproducibility and Ablation Addendum\n\n"
        "This note accompanies `titan_fagbi_language_edited_clean_repro_ablation.docx`.\n\n"
        "- Added Table M1 with electrostatics, PDB2PQR, surface, solver, preconditioner, audit, hardware, timer, and RSS parameters.\n"
        "- Added Results subsection 3.5 on ablation and sensitivity evidence.\n"
        "- Clarified that Direct TABI means direct dense panel-panel summation of the TABI boundary-integral operator.\n"
        "- Added ILU(0) setup-overhead wording using profile timing fields from the manuscript-facing rows.\n",
        encoding="utf-8",
    )
    return md


def write_qa(rows):
    with QA_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["check", "status", "detail"])
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)
    doc = Document(str(INPUT))
    add_direct_tabi_sentence(doc)
    add_methods_table(doc)
    add_results_sensitivity(doc)
    doc.save(str(OUTPUT))
    md = save_markdown_note()

    reloaded = Document(str(OUTPUT))
    text = "\n".join(p.text for p in reloaded.paragraphs)
    rows = [
        {
            "check": "direct_tabi_definition",
            "status": "pass" if "Direct TABI denotes direct summation" in text else "fail",
            "detail": "dense panel-panel definition present",
        },
        {
            "check": "reproducibility_table",
            "status": "pass" if "Table M1. Reproducibility parameters" in text else "fail",
            "detail": "Methods reproducibility table caption present",
        },
        {
            "check": "ablation_section",
            "status": "pass" if "3.5. Ablation and Sensitivity" in text else "fail",
            "detail": "Results ablation section present",
        },
        {
            "check": "output_docx",
            "status": "pass" if OUTPUT.exists() and OUTPUT.stat().st_size > 0 else "fail",
            "detail": str(OUTPUT),
        },
        {
            "check": "markdown_note",
            "status": "pass" if md.exists() else "fail",
            "detail": str(md),
        },
    ]
    write_qa(rows)
    failed = [r for r in rows if r["status"] != "pass"]
    if failed:
        raise RuntimeError(f"QA failed: {failed}")
    print(f"wrote {OUTPUT}")
    print(f"qa {QA_CSV}")


if __name__ == "__main__":
    main()
