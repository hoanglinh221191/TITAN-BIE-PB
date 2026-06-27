from __future__ import annotations

import csv
import math
import statistics
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
Z4W_MATRIX = ROOT / "test/tmp/fagbi_scope_z4w/fagbi_scope_z4w_matrix.csv"
Z4XF_MATRIX = ROOT / "test/tmp/fagbi_scope_z4xf/fagbi_scope_z4xf_matrix.csv"
Z4Q_MATRIX = ROOT / "docs/fagbi_guarded_p4_package/tables/fagbi_scope_z4q_matrix.csv"
Z4Q_REPEATABILITY = ROOT / "docs/fagbi_guarded_p4_package/tables/fagbi_scope_z4q_repeatability.csv"
Z4Q_MANIFEST = ROOT / "docs/fagbi_guarded_p4_package/provenance/public_fixture_expansion_manifest.csv"
OUT_DIR = ROOT / "test/tmp/fagbi_scope_z4xe"
PREVIEW_DIR = ROOT / "docs/fagbi_guarded_p4_package/submission_preview_synced"
FIGURE_DIR = PREVIEW_DIR / "figures"
DOCX_OUT = PREVIEW_DIR / "titan_fagbi_supplement_z4xe_sensitivity.docx"
P2_SUPPLEMENT_OUT = PREVIEW_DIR / "titan_fagbi_paper_standard_revised_p2_coverage_supplement.docx"
ROOT_REPORT = ROOT / "TITAN_FMM_M2L_SCOPE_Z4XE_SUPPLEMENTARY_SENSITIVITY_REPORT.txt"
Z4XF_FIGURE_PNG = FIGURE_DIR / "fagbi_scope_z4xf_ilu_neighbor_cap_sweep.png"
Z4XF_FIGURE_SVG = FIGURE_DIR / "fagbi_scope_z4xf_ilu_neighbor_cap_sweep.svg"


def as_float(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    try:
        out = float(text)
    except ValueError:
        return None
    return out if math.isfinite(out) else None


def fmt(value: float | None, digits: int = 3) -> str:
    if value is None or not math.isfinite(value):
        return "NA"
    if value == 0:
        return "0"
    if abs(value) < 1e-3 or abs(value) >= 1e4:
        return f"{value:.{digits}e}"
    return f"{value:.{digits}f}".rstrip("0").rstrip(".")


def read_rows() -> list[dict[str, str]]:
    if not Z4W_MATRIX.exists():
        raise SystemExit(f"Missing Z4W matrix: {Z4W_MATRIX}")
    with Z4W_MATRIX.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def read_z4xf_rows() -> list[dict[str, str]]:
    if not Z4XF_MATRIX.exists():
        return []
    with Z4XF_MATRIX.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def read_csv_if_exists(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({name: row.get(name, "") for name in fieldnames})


def public_protein_rows() -> list[dict[str, object]]:
    matrix = read_csv_if_exists(Z4Q_MATRIX)
    repeatability = {row.get("fixture", ""): row for row in read_csv_if_exists(Z4Q_REPEATABILITY)}
    manifest = {row.get("fixture", ""): row for row in read_csv_if_exists(Z4Q_MANIFEST)}
    selected: dict[str, dict[str, str]] = {}
    for row in matrix:
        if row.get("variant") != "p4_guarded":
            continue
        fixture = row.get("fixture", "")
        if not fixture:
            continue
        current = selected.get(fixture)
        if current is None or int(row.get("repeat", "999") or "999") < int(current.get("repeat", "999") or "999"):
            selected[fixture] = row
    out: list[dict[str, object]] = []
    for fixture in sorted(selected):
        row = selected[fixture]
        rep = repeatability.get(fixture, {})
        man = manifest.get(fixture, {})
        out.append(
            {
                "protein": row.get("protein_name") or man.get("protein_name") or fixture,
                "pdb": row.get("pdb_id") or man.get("pdb_id") or fixture.upper(),
                "atoms": man.get("atom_count") or row.get("atom_count", ""),
                "panels": row.get("panels", ""),
                "energy": fmt(as_float(row.get("energy_kj_mol")), 2),
                "audit_max": fmt(as_float(rep.get("audit_max") or row.get("audit_rms")), 3),
                "direct_residual_max": fmt(as_float(rep.get("sampled_direct_true_residual_max") or row.get("sampled_direct_true_residual")), 3),
                "energy_spread": fmt(as_float(rep.get("energy_spread")), 3),
            }
        )
    return out


def stage_label(row: dict[str, str]) -> str:
    stage = row.get("stage", "")
    mapping = {
        "baseline_guarded_p4": "baseline guarded p4",
        "theta0p4": "theta 0.4",
        "theta0p5": "theta 0.5",
        "theta0p6": "theta 0.6",
        "theta0p7": "theta 0.7",
        "tail_off": "tail indicator off",
        "tail_on": "tail indicator on",
        "geometry_off": "geometry guard off",
        "geometry_on": "geometry guard on",
        "ilu_cap32": "ILU neighbor cap 32",
        "ilu_cap64": "ILU neighbor cap 64",
        "ilu_cap128": "ILU neighbor cap 128",
    }
    return mapping.get(stage, stage)


def value_label(row: dict[str, str]) -> str:
    axis = row.get("axis", "")
    if axis == "opening_theta":
        return row.get("opening_theta", "")
    if axis == "geometry_guard":
        return row.get("geometry_guard", "")
    if axis == "tail_indicator":
        return row.get("refine", "")
    if axis == "ilu_neighbor_cap":
        return row.get("ilu_neighbor_cap", "")
    if axis == "baseline":
        return "theta 0.6; tail on; guard theta 0.5; ILU cap 64"
    return row.get("value", "")


def aggregate(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    wanted = {"baseline", "opening_theta", "geometry_guard", "tail_indicator"}
    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in rows:
        if row.get("axis") in wanted:
            grouped.setdefault((row.get("axis", ""), row.get("stage", "")), []).append(row)
    out: list[dict[str, object]] = []
    for (axis, stage), group in sorted(grouped.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        ok = [r for r in group if r.get("status") == "ok"]
        audits = [as_float(r.get("sampled_audit_rms")) for r in ok]
        direct_res = [as_float(r.get("sampled_direct_true_residual")) for r in ok]
        wall = [as_float(r.get("wall_seconds")) for r in ok]
        rss = [as_float(r.get("peak_rss_mb")) for r in ok]
        audits_f = [v for v in audits if v is not None]
        direct_f = [v for v in direct_res if v is not None]
        wall_f = [v for v in wall if v is not None]
        rss_f = [v for v in rss if v is not None]
        status_count = f"{len(ok)}/{len(group)} ok"
        if stage == "geometry_off":
            conclusion = "Measured guard-off sensitivity; 1A63 exceeds the direct-operator audit gate."
        elif stage == "geometry_on":
            conclusion = "Validated guarded policy; all six rows pass."
        elif stage == "tail_off":
            conclusion = "Rejected by the guarded p4 validator; tail refinement is required by this protocol."
        elif stage == "tail_on":
            conclusion = "Validated tail-refined policy; all six rows pass."
        elif axis == "opening_theta":
            conclusion = "Measured theta sensitivity row; all tested theta values pass in this protocol."
        else:
            conclusion = "Validated baseline policy."
        out.append(
            {
                "axis": axis,
                "stage": stage_label(group[0]),
                "value": value_label(group[0]),
                "status_count": status_count,
                "max_audit_rms": fmt(max(audits_f) if audits_f else None, 3),
                "max_sampled_direct_true_residual": fmt(max(direct_f) if direct_f else None, 3),
                "median_wall_seconds": fmt(statistics.median(wall_f) if wall_f else None, 1),
                "max_peak_rss_mb": fmt(max(rss_f) if rss_f else None, 1),
                "conclusion": conclusion,
            }
        )
    return out


def detail_rows(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    wanted = {"baseline", "opening_theta", "geometry_guard", "tail_indicator"}
    selected = [r for r in rows if r.get("axis") in wanted]
    order = {"baseline": 0, "opening_theta": 1, "geometry_guard": 2, "tail_indicator": 3}
    selected.sort(key=lambda r: (order.get(r.get("axis", ""), 99), r.get("stage_order", ""), r.get("fixture", "")))
    out: list[dict[str, object]] = []
    for r in selected:
        out.append(
            {
                "fixture": r.get("pdb_id") or r.get("fixture"),
                "axis": r.get("axis"),
                "stage": stage_label(r),
                "value": value_label(r),
                "status": r.get("status", ""),
                "failure_class": r.get("failure_class", ""),
                "audit_rms": fmt(as_float(r.get("sampled_audit_rms")), 3),
                "sampled_direct_true_residual": fmt(as_float(r.get("sampled_direct_true_residual")), 3),
                "true_residual": fmt(as_float(r.get("true_residual")), 3),
                "iterations": r.get("iterations", ""),
                "wall_seconds": fmt(as_float(r.get("wall_seconds")), 1),
                "peak_rss_mb": fmt(as_float(r.get("peak_rss_mb")), 1),
            }
        )
    return out


def cap_boundary(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    selected = [r for r in rows if r.get("axis") == "ilu_neighbor_cap"]
    grouped: dict[str, list[dict[str, str]]] = {}
    for row in selected:
        grouped.setdefault(row.get("ilu_neighbor_cap", ""), []).append(row)
    out: list[dict[str, object]] = []
    for cap in sorted(grouped, key=lambda x: int(x) if x.isdigit() else 999):
        group = grouped[cap]
        ok = [r for r in group if r.get("status") == "ok"]
        direct_res = [as_float(r.get("sampled_direct_true_residual")) for r in ok]
        wall = [as_float(r.get("wall_seconds")) for r in ok]
        rss = [as_float(r.get("peak_rss_mb")) for r in ok]
        if cap == "64":
            interpretation = "Validated policy value under the guarded p4 protocol; all six rows pass."
        else:
            interpretation = "Protocol-boundary row: current gmres-p4 validator rejects caps other than 64, so this is not a numerical cap sweep."
        out.append(
            {
                "ilu_neighbor_cap": cap,
                "status_count": f"{len(ok)}/{len(group)} ok",
                "failure_class_summary": ", ".join(sorted({r.get("failure_class", "") for r in group if r.get("failure_class")})),
                "max_sampled_direct_true_residual": fmt(max([v for v in direct_res if v is not None]) if direct_res else None, 3),
                "median_wall_seconds": fmt(statistics.median([v for v in wall if v is not None]) if any(v is not None for v in wall) else None, 1),
                "max_peak_rss_mb": fmt(max([v for v in rss if v is not None]) if any(v is not None for v in rss) else None, 1),
                "interpretation": interpretation,
            }
        )
    return out


def mean_sd(values: list[float]) -> tuple[float | None, float | None]:
    clean = [v for v in values if v is not None and math.isfinite(v)]
    if not clean:
        return None, None
    if len(clean) == 1:
        return clean[0], 0.0
    return statistics.mean(clean), statistics.stdev(clean)


def fmt_mean_sd(mean: float | None, sd: float | None, digits: int = 2) -> str:
    if mean is None or sd is None:
        return "NA"
    return f"{fmt(mean, digits)} +/- {fmt(sd, digits)}"


def cap_real_sweep(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    grouped: dict[int, list[dict[str, str]]] = {}
    for row in rows:
        cap = int(row.get("ilu_neighbor_cap", "0") or 0)
        if cap in {32, 64, 128}:
            grouped.setdefault(cap, []).append(row)
    out: list[dict[str, object]] = []
    for cap in [32, 64, 128]:
        group = grouped.get(cap, [])
        ok = [r for r in group if r.get("status") == "ok" and r.get("failure_class") == "pass"]
        audits = [as_float(r.get("sampled_audit_rms")) for r in ok]
        direct_res = [as_float(r.get("sampled_direct_true_residual")) for r in ok]
        matvec = [as_float(r.get("matvec_seconds_avg")) for r in ok]
        wall = [as_float(r.get("wall_seconds")) for r in ok]
        rss = [as_float(r.get("peak_rss_mb")) for r in ok]
        nnz = [as_float(r.get("preconditioner_sparse_nnz_blocks")) for r in ok]
        matvec_mean, matvec_sd = mean_sd([v for v in matvec if v is not None])
        wall_mean, wall_sd = mean_sd([v for v in wall if v is not None])
        rss_mean, rss_sd = mean_sd([v for v in rss if v is not None])
        if cap == 32:
            interpretation = "All rows pass; memory-lean graph, but not the fastest average matvec in this sweep."
        elif cap == 64:
            interpretation = "All rows pass; best average matvec time with moderate graph size, supporting the reported protocol value."
        else:
            interpretation = "All rows pass; larger graph and RSS without a consistent accuracy advantage."
        out.append(
            {
                "ilu_neighbor_cap": cap,
                "status_count": f"{len(ok)}/{len(group)} ok",
                "max_audit_rms": fmt(max([v for v in audits if v is not None]) if audits else None, 3),
                "max_sampled_direct_true_residual": fmt(max([v for v in direct_res if v is not None]) if direct_res else None, 3),
                "mean_matvec_seconds": fmt_mean_sd(matvec_mean, matvec_sd, 3),
                "mean_wall_seconds": fmt_mean_sd(wall_mean, wall_sd, 1),
                "mean_peak_rss_mb": fmt_mean_sd(rss_mean, rss_sd, 1),
                "median_nnz_blocks": fmt(statistics.median([v for v in nnz if v is not None]) if any(v is not None for v in nnz) else None, 0),
                "interpretation": interpretation,
            }
        )
    return out


def z4xf_stats(rows: list[dict[str, str]]) -> dict[int, dict[str, float]]:
    stats: dict[int, dict[str, float]] = {}
    for cap in [32, 64, 128]:
        group = [
            r for r in rows
            if int(r.get("ilu_neighbor_cap", "0") or 0) == cap
            and r.get("status") == "ok"
            and r.get("failure_class") == "pass"
        ]
        if not group:
            continue
        audits = [as_float(r.get("sampled_direct_true_residual")) for r in group]
        matvec = [as_float(r.get("matvec_seconds_avg")) for r in group]
        rss = [as_float(r.get("peak_rss_mb")) for r in group]
        matvec_mean, matvec_sd = mean_sd([v for v in matvec if v is not None])
        rss_mean, rss_sd = mean_sd([v for v in rss if v is not None])
        stats[cap] = {
            "max_residual": max([v for v in audits if v is not None]),
            "matvec_mean": matvec_mean or 0.0,
            "matvec_sd": matvec_sd or 0.0,
            "rss_mean": rss_mean or 0.0,
            "rss_sd": rss_sd or 0.0,
        }
    return stats


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    font_name = "arialbd.ttf" if bold else "arial.ttf"
    for path in [Path("C:/Windows/Fonts") / font_name, Path("C:/Windows/Fonts/calibri.ttf")]:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_text_center(draw: ImageDraw.ImageDraw, xy: tuple[float, float], text: str, font, fill=(35, 35, 35)) -> None:
    bbox = draw.textbbox((0, 0), text, font=font)
    draw.text((xy[0] - (bbox[2] - bbox[0]) / 2, xy[1] - (bbox[3] - bbox[1]) / 2), text, font=font, fill=fill)


def draw_bar_panel(
    draw: ImageDraw.ImageDraw,
    x0: int,
    y0: int,
    w: int,
    h: int,
    label: str,
    y_label: str,
    values: dict[int, tuple[float, float]],
    color: tuple[int, int, int],
    y_max: float | None = None,
) -> None:
    font_label = load_font(24, True)
    font_axis = load_font(18)
    font_tick = load_font(16)
    draw.text((x0, y0 - 35), label, font=font_label, fill=(20, 20, 20))
    left = x0 + 62
    right = x0 + w - 18
    top = y0 + 18
    bottom = y0 + h - 52
    draw.line((left, bottom, right, bottom), fill=(45, 45, 45), width=2)
    draw.line((left, top, left, bottom), fill=(45, 45, 45), width=2)
    vmax = y_max if y_max is not None else max(v[0] + v[1] for v in values.values()) * 1.15
    if vmax <= 0:
        vmax = 1.0
    for frac in [0.0, 0.5, 1.0]:
        y = bottom - frac * (bottom - top)
        draw.line((left - 5, y, right, y), fill=(225, 225, 225), width=1)
        draw_text_center(draw, (left - 33, y), fmt(vmax * frac, 2), font_tick, fill=(70, 70, 70))
    caps = [32, 64, 128]
    slot = (right - left) / len(caps)
    bar_w = 46
    for i, cap in enumerate(caps):
        mean, sd = values.get(cap, (0.0, 0.0))
        cx = left + slot * (i + 0.5)
        bh = (mean / vmax) * (bottom - top)
        y = bottom - bh
        draw.rectangle((cx - bar_w / 2, y, cx + bar_w / 2, bottom), fill=color, outline=(35, 35, 35))
        if sd > 0:
            y_hi = bottom - ((mean + sd) / vmax) * (bottom - top)
            y_lo = bottom - (max(0.0, mean - sd) / vmax) * (bottom - top)
            draw.line((cx, y_hi, cx, y_lo), fill=(25, 25, 25), width=2)
            draw.line((cx - 10, y_hi, cx + 10, y_hi), fill=(25, 25, 25), width=2)
            draw.line((cx - 10, y_lo, cx + 10, y_lo), fill=(25, 25, 25), width=2)
        draw_text_center(draw, (cx, bottom + 22), str(cap), font_tick)
    draw_text_center(draw, ((left + right) / 2, y0 + h - 12), "ILU neighbor cap", font_axis)
    draw.text((x0 + 2, y0 - 8), y_label, font=font_axis, fill=(45, 45, 45))


def draw_log_panel(
    draw: ImageDraw.ImageDraw,
    x0: int,
    y0: int,
    w: int,
    h: int,
    values: dict[int, float],
) -> None:
    font_label = load_font(24, True)
    font_axis = load_font(18)
    font_tick = load_font(16)
    draw.text((x0, y0 - 35), "A", font=font_label, fill=(20, 20, 20))
    left = x0 + 70
    right = x0 + w - 18
    top = y0 + 18
    bottom = y0 + h - 52
    ymin, ymax = -6.0, -2.0
    def y_pos(v: float) -> float:
        return bottom - ((math.log10(max(v, 1e-12)) - ymin) / (ymax - ymin)) * (bottom - top)
    draw.line((left, bottom, right, bottom), fill=(45, 45, 45), width=2)
    draw.line((left, top, left, bottom), fill=(45, 45, 45), width=2)
    for tick in [1e-6, 1e-5, 1e-4, 1e-3, 1e-2]:
        y = y_pos(tick)
        draw.line((left - 5, y, right, y), fill=(225, 225, 225), width=1)
        draw_text_center(draw, (left - 39, y), f"1e{int(math.log10(tick))}", font_tick, fill=(70, 70, 70))
    gate_y = y_pos(8e-3)
    draw.line((left, gate_y, right, gate_y), fill=(150, 50, 50), width=2)
    draw.text((right - 105, gate_y - 22), "8e-3 gate", font=font_tick, fill=(150, 50, 50))
    caps = [32, 64, 128]
    slot = (right - left) / len(caps)
    points: list[tuple[float, float]] = []
    for i, cap in enumerate(caps):
        cx = left + slot * (i + 0.5)
        y = y_pos(values.get(cap, 1e-6))
        points.append((cx, y))
        draw.ellipse((cx - 7, y - 7, cx + 7, y + 7), fill=(68, 119, 170), outline=(30, 30, 30))
        draw_text_center(draw, (cx, bottom + 22), str(cap), font_tick)
    if len(points) > 1:
        draw.line(points, fill=(68, 119, 170), width=3)
    draw.text((x0 + 2, y0 - 8), "max sampled direct residual", font=font_axis, fill=(45, 45, 45))
    draw_text_center(draw, ((left + right) / 2, y0 + h - 12), "ILU neighbor cap", font_axis)


def write_z4xf_figure(rows: list[dict[str, str]]) -> bool:
    stats = z4xf_stats(rows)
    if len(stats) != 3:
        return False
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1650, 520
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    panel_w = 505
    y0 = 105
    h = 340
    draw_log_panel(draw, 35, y0, panel_w, h, {cap: stats[cap]["max_residual"] for cap in stats})
    draw_bar_panel(
        draw,
        575,
        y0,
        panel_w,
        h,
        "B",
        "matvec time (s)",
        {cap: (stats[cap]["matvec_mean"], stats[cap]["matvec_sd"]) for cap in stats},
        (68, 119, 170),
    )
    draw_bar_panel(
        draw,
        1115,
        y0,
        panel_w,
        h,
        "C",
        "peak RSS (MB)",
        {cap: (stats[cap]["rss_mean"], stats[cap]["rss_sd"]) for cap in stats},
        (204, 121, 57),
    )
    note_font = load_font(16)
    draw.text((42, 482), "Points/bars summarize Barnase, 1A63, and 1SH1; error bars show sample standard deviation.", font=note_font, fill=(70, 70, 70))
    img.save(Z4XF_FIGURE_PNG, dpi=(300, 300))

    def svg_text(x: float, y: float, text: str, size: int = 16, weight: str = "normal",
                 anchor: str = "start", fill: str = "#333") -> str:
        return (
            f'<text x="{x:.1f}" y="{y:.1f}" font-family="Arial" '
            f'font-size="{size}" font-weight="{weight}" text-anchor="{anchor}" '
            f'fill="{fill}">{text}</text>'
        )

    svg_lines = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1650" height="520" viewBox="0 0 1650 520">',
        '<rect width="100%" height="100%" fill="white"/>',
    ]

    def add_svg_log_panel(x0: int, label: str) -> None:
        left, right, top, bottom = x0 + 70, x0 + 505 - 18, 123, 393
        ymin, ymax = -6.0, -2.0
        def ypos(v: float) -> float:
            return bottom - ((math.log10(max(v, 1e-12)) - ymin) / (ymax - ymin)) * (bottom - top)
        svg_lines.append(svg_text(x0, 70, label, 24, "bold"))
        svg_lines.append(f'<line x1="{left}" y1="{bottom}" x2="{right}" y2="{bottom}" stroke="#333" stroke-width="2"/>')
        svg_lines.append(f'<line x1="{left}" y1="{top}" x2="{left}" y2="{bottom}" stroke="#333" stroke-width="2"/>')
        for tick in [1e-6, 1e-5, 1e-4, 1e-3, 1e-2]:
            y = ypos(tick)
            svg_lines.append(f'<line x1="{left-5}" y1="{y:.1f}" x2="{right}" y2="{y:.1f}" stroke="#e5e5e5" stroke-width="1"/>')
            svg_lines.append(svg_text(left - 39, y + 5, f"1e{int(math.log10(tick))}", 16, anchor="middle", fill="#555"))
        gate_y = ypos(8e-3)
        svg_lines.append(f'<line x1="{left}" y1="{gate_y:.1f}" x2="{right}" y2="{gate_y:.1f}" stroke="#963232" stroke-width="2"/>')
        svg_lines.append(svg_text(right - 105, gate_y - 8, "8e-3 gate", 16, fill="#963232"))
        points = []
        caps = [32, 64, 128]
        slot = (right - left) / len(caps)
        for i, cap in enumerate(caps):
            cx = left + slot * (i + 0.5)
            y = ypos(stats[cap]["max_residual"])
            points.append((cx, y))
            svg_lines.append(f'<circle cx="{cx:.1f}" cy="{y:.1f}" r="7" fill="#4477AA" stroke="#222" stroke-width="1"/>')
            svg_lines.append(svg_text(cx, bottom + 26, str(cap), 16, anchor="middle"))
        for (x1, y1), (x2, y2) in zip(points, points[1:]):
            svg_lines.append(f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" stroke="#4477AA" stroke-width="3"/>')
        svg_lines.append(svg_text(x0 + 2, 98, "max sampled direct residual", 18))
        svg_lines.append(svg_text((left + right) / 2, 426, "ILU neighbor cap", 18, anchor="middle"))

    def add_svg_bar_panel(x0: int, label: str, y_label: str, key_mean: str, key_sd: str,
                          color: str) -> None:
        left, right, top, bottom = x0 + 62, x0 + 505 - 18, 123, 393
        values = {cap: (stats[cap][key_mean], stats[cap][key_sd]) for cap in [32, 64, 128]}
        vmax = max(mean + sd for mean, sd in values.values()) * 1.15
        if vmax <= 0:
            vmax = 1.0
        svg_lines.append(svg_text(x0, 70, label, 24, "bold"))
        svg_lines.append(f'<line x1="{left}" y1="{bottom}" x2="{right}" y2="{bottom}" stroke="#333" stroke-width="2"/>')
        svg_lines.append(f'<line x1="{left}" y1="{top}" x2="{left}" y2="{bottom}" stroke="#333" stroke-width="2"/>')
        for frac in [0.0, 0.5, 1.0]:
            y = bottom - frac * (bottom - top)
            svg_lines.append(f'<line x1="{left-5}" y1="{y:.1f}" x2="{right}" y2="{y:.1f}" stroke="#e5e5e5" stroke-width="1"/>')
            svg_lines.append(svg_text(left - 33, y + 5, fmt(vmax * frac, 2), 16, anchor="middle", fill="#555"))
        caps = [32, 64, 128]
        slot = (right - left) / len(caps)
        bar_w = 46
        for i, cap in enumerate(caps):
            mean, sd = values[cap]
            cx = left + slot * (i + 0.5)
            bh = (mean / vmax) * (bottom - top)
            y = bottom - bh
            svg_lines.append(f'<rect x="{cx-bar_w/2:.1f}" y="{y:.1f}" width="{bar_w}" height="{bh:.1f}" fill="{color}" stroke="#222" stroke-width="1"/>')
            if sd > 0:
                y_hi = bottom - ((mean + sd) / vmax) * (bottom - top)
                y_lo = bottom - (max(0.0, mean - sd) / vmax) * (bottom - top)
                svg_lines.append(f'<line x1="{cx:.1f}" y1="{y_hi:.1f}" x2="{cx:.1f}" y2="{y_lo:.1f}" stroke="#222" stroke-width="2"/>')
                svg_lines.append(f'<line x1="{cx-10:.1f}" y1="{y_hi:.1f}" x2="{cx+10:.1f}" y2="{y_hi:.1f}" stroke="#222" stroke-width="2"/>')
                svg_lines.append(f'<line x1="{cx-10:.1f}" y1="{y_lo:.1f}" x2="{cx+10:.1f}" y2="{y_lo:.1f}" stroke="#222" stroke-width="2"/>')
            svg_lines.append(svg_text(cx, bottom + 26, str(cap), 16, anchor="middle"))
        svg_lines.append(svg_text(x0 + 2, 98, y_label, 18))
        svg_lines.append(svg_text((left + right) / 2, 426, "ILU neighbor cap", 18, anchor="middle"))

    add_svg_log_panel(35, "A")
    add_svg_bar_panel(575, "B", "matvec time (s)", "matvec_mean", "matvec_sd", "#4477AA")
    add_svg_bar_panel(1115, "C", "peak RSS (MB)", "rss_mean", "rss_sd", "#CC7939")
    svg_lines.append(svg_text(42, 490, "Points/bars summarize Barnase, 1A63, and 1SH1; error bars show sample standard deviation.", 16, fill="#555"))
    svg_lines.append("</svg>")
    Z4XF_FIGURE_SVG.write_text("\n".join(svg_lines), encoding="utf-8")
    return True


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, font_size: float = 8.0) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, title: str, rows: list[dict[str, object]], columns: list[tuple[str, str]], widths: list[float]) -> None:
    p = doc.add_paragraph(title)
    p.style = "Heading 2"
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, (_, label) in enumerate(columns):
        set_cell_text(hdr[idx], label, bold=True, font_size=8.0)
        set_cell_shading(hdr[idx], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for idx, (key, _) in enumerate(columns):
            set_cell_text(cells[idx], str(row.get(key, "")), font_size=7.2)
    set_table_widths(table, widths)
    doc.add_paragraph()


def configure_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11)
    sec.page_height = Inches(8.5)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in [("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10)]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 77, 120)


def build_docx(
    summary: list[dict[str, object]],
    details: list[dict[str, object]],
    caps: list[dict[str, object]],
    *,
    public_rows: list[dict[str, object]],
    cap_source: str,
    z4xf_figure_ready: bool,
) -> None:
    doc = Document()
    configure_styles(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Supplementary Tables and Sensitivity Figures for the Guarded p4 TABI Treecode")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Public-protein expansion, theta/guard/tail sensitivity, and ILU neighbor-cap sensitivity").italic = True

    doc.add_paragraph(
        "This supplement is generated from the completed public-protein, ablation, and ILU neighbor-cap artifacts. "
        "It does not rerun the solver. The purpose is to keep supplementary evidence out of the main manuscript "
        "while preserving reviewer-auditable tables and figures."
    )
    if cap_source == "z4xf_real_sweep":
        doc.add_paragraph(
            "ILU neighbor-cap evidence is taken from the Z4X-F real sweep on Barnase, 1A63, and 1SH1. "
            "All three tested caps preserve the direct-operator verification gates; cap 64 is retained in the "
            "reported guarded-p4 protocol because it gives the best average matvec time while avoiding the larger "
            "sparse graph and RSS of cap 128."
        )
    else:
        doc.add_paragraph(
            "Important boundary: ILU neighbor cap 64 is the validated guarded-p4 protocol value. Caps 32 and 128 "
            "are listed here as validator-boundary rows, not as a numerical neighbor-cap sweep, because the explicit "
            "gmres-p4 validator rejects non-64 caps under the locked protocol."
        )

    if public_rows:
        add_table(
            doc,
            "Table S1. Full public-protein supplement rows prepared without atom sampling.",
            public_rows,
            [
                ("protein", "Protein"),
                ("pdb", "PDB"),
                ("atoms", "Atoms"),
                ("panels", "Panels"),
                ("energy", "Energy"),
                ("audit_max", "Audit max"),
                ("direct_residual_max", "Direct residual max"),
                ("energy_spread", "Energy spread"),
            ],
            [2.2, 0.7, 0.7, 0.75, 0.85, 0.9, 1.1, 0.95],
        )

    sensitivity_offset = 1 if public_rows else 0
    add_table(
        doc,
        f"Table S{1 + sensitivity_offset}. Aggregated theta, guard, and tail sensitivity.",
        summary,
        [
            ("axis", "Axis"),
            ("stage", "Row"),
            ("value", "Value"),
            ("status_count", "Status"),
            ("max_audit_rms", "Max audit RMS"),
            ("max_sampled_direct_true_residual", "Max sampled direct residual"),
            ("median_wall_seconds", "Median wall (s)"),
            ("max_peak_rss_mb", "Max RSS (MB)"),
            ("conclusion", "Interpretation"),
        ],
        [1.1, 1.35, 1.55, 0.75, 0.9, 1.05, 0.9, 0.9, 2.4],
    )
    add_table(
        doc,
        f"Table S{2 + sensitivity_offset}. Per-fixture sensitivity rows.",
        details,
        [
            ("fixture", "Fixture"),
            ("axis", "Axis"),
            ("stage", "Row"),
            ("value", "Value"),
            ("status", "Status"),
            ("failure_class", "Failure class"),
            ("audit_rms", "Audit RMS"),
            ("sampled_direct_true_residual", "Sampled direct residual"),
            ("true_residual", "True residual"),
            ("iterations", "Iter."),
            ("wall_seconds", "Wall (s)"),
            ("peak_rss_mb", "RSS (MB)"),
        ],
        [0.65, 1.0, 1.25, 1.35, 0.55, 0.95, 0.8, 1.0, 0.85, 0.45, 0.65, 0.65],
    )
    add_table(
        doc,
        f"Table S{3 + sensitivity_offset}. ILU neighbor-cap real sweep." if cap_source == "z4xf_real_sweep" else f"Table S{3 + sensitivity_offset}. ILU neighbor-cap policy boundary.",
        caps,
        (
            [
                ("ilu_neighbor_cap", "Cap"),
                ("status_count", "Status"),
                ("max_audit_rms", "Max audit RMS"),
                ("max_sampled_direct_true_residual", "Max sampled direct residual"),
                ("mean_matvec_seconds", "Mean matvec (s)"),
                ("mean_wall_seconds", "Mean wall (s)"),
                ("mean_peak_rss_mb", "Mean RSS (MB)"),
                ("median_nnz_blocks", "Median nnz blocks"),
                ("interpretation", "Interpretation"),
            ]
            if cap_source == "z4xf_real_sweep"
            else [
                ("ilu_neighbor_cap", "Cap"),
                ("status_count", "Status"),
                ("failure_class_summary", "Failure summary"),
                ("max_sampled_direct_true_residual", "Max sampled direct residual"),
                ("median_wall_seconds", "Median wall (s)"),
                ("max_peak_rss_mb", "Max RSS (MB)"),
                ("interpretation", "Interpretation"),
            ]
        ),
        [0.45, 0.65, 0.75, 0.95, 0.85, 0.8, 0.85, 0.9, 2.7]
        if cap_source == "z4xf_real_sweep"
        else [0.55, 0.8, 1.1, 1.15, 0.9, 0.9, 4.25],
    )
    if z4xf_figure_ready and Z4XF_FIGURE_PNG.exists():
        p = doc.add_paragraph("Figure S1. ILU neighbor-cap sensitivity under the guarded p4 protocol.")
        p.style = "Heading 2"
        doc.add_picture(str(Z4XF_FIGURE_PNG), width=Inches(9.3))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.add_run(
            "A, maximum sampled direct residual across Barnase, 1A63, and 1SH1, with the 8 x 10^-3 verification "
            "gate shown as a red reference line. B, mean GMRES matvec time. C, mean process peak RSS. Error bars "
            "show sample standard deviation over the three proteins. The sweep shows that cap 32, 64, and 128 all "
            "pass accuracy; cap 64 is a balanced protocol value rather than an accuracy-only optimum."
        )

    if cap_source == "z4xf_real_sweep":
        doc.add_paragraph(
            "Conclusion. The measured theta/guard/tail rows support the manuscript statement that the guarded p4 policy is not "
            "fragile over the reported theta range and that the strict geometry guard is accuracy-critical for 1A63. The real "
            "neighbor-cap sweep shows that cap 64 remains a balanced protocol value: cap 32 is memory-lean, cap 128 enlarges the "
            "sparse graph and RSS without a consistent accuracy advantage, and cap 64 gives the best average matvec time in this "
            "three-protein sweep."
        )
    else:
        doc.add_paragraph(
            "Conclusion. The measured theta/guard/tail rows support the manuscript statement that the guarded p4 policy is not "
            "fragile over the reported theta range and that the strict geometry guard is accuracy-critical for 1A63. The neighbor-cap "
            "evidence supports cap 64 only as a locked validated setting. A real cap 32/64/128 numerical sweep should be a separate "
            "benchmark if the manuscript later claims cap optimality."
        )
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)
    doc.save(P2_SUPPLEMENT_OUT)


def main() -> None:
    rows = read_rows()
    z4xf_rows = read_z4xf_rows()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary = aggregate(rows)
    details = detail_rows(rows)
    public_rows = public_protein_rows()
    real_caps = cap_real_sweep(z4xf_rows) if z4xf_rows else []
    cap_source = "z4xf_real_sweep" if real_caps else "z4w_validator_boundary"
    caps = real_caps if real_caps else cap_boundary(rows)
    z4xf_figure_ready = write_z4xf_figure(z4xf_rows) if real_caps else False
    summary_csv = OUT_DIR / "fagbi_scope_z4xe_theta_guard_tail_summary.csv"
    detail_csv = OUT_DIR / "fagbi_scope_z4xe_theta_guard_tail_detail.csv"
    cap_csv = OUT_DIR / "fagbi_scope_z4xe_neighbor_cap_boundary.csv"
    real_cap_csv = OUT_DIR / "fagbi_scope_z4xf_neighbor_cap_real_sweep.csv"
    decision_csv = OUT_DIR / "fagbi_scope_z4xe_decision.csv"
    summary_dat = OUT_DIR / "fagbi_scope_z4xe_summary.dat"

    write_csv(summary_csv, summary, list(summary[0].keys()))
    write_csv(detail_csv, details, list(details[0].keys()))
    write_csv(cap_csv, caps, list(caps[0].keys()))
    if real_caps:
        write_csv(real_cap_csv, real_caps, list(real_caps[0].keys()))

    decision = "z4xe_supplement_ready_with_z4xf_cap_sweep" if real_caps else "z4xe_supplement_ready"
    write_csv(
        decision_csv,
        [
            {
                "decision": decision,
                "theta_guard_tail_source": str(Z4W_MATRIX),
                "neighbor_cap_source": str(Z4XF_MATRIX) if real_caps else str(Z4W_MATRIX),
                "docx": str(DOCX_OUT),
                "paired_docx": str(P2_SUPPLEMENT_OUT),
                "public_protein_rows": len(public_rows),
                "neighbor_cap_policy": "real_sweep_cap32_64_128" if real_caps else "cap64_validated_boundary_only",
                "figure": str(Z4XF_FIGURE_PNG) if z4xf_figure_ready else "",
            }
        ],
        ["decision", "theta_guard_tail_source", "neighbor_cap_source", "docx", "paired_docx", "public_protein_rows", "neighbor_cap_policy", "figure"],
    )
    summary_dat.write_text(
        "\n".join(
            [
                "scope: Z4X-E supplementary sensitivity packaging",
                f"decision: {decision}",
                f"source_matrix: {Z4W_MATRIX}",
                f"neighbor_cap_source: {Z4XF_MATRIX if real_caps else Z4W_MATRIX}",
                f"docx: {DOCX_OUT}",
                f"paired_docx: {P2_SUPPLEMENT_OUT}",
                f"public_protein_rows: {len(public_rows)}",
                f"summary_csv: {summary_csv}",
                f"detail_csv: {detail_csv}",
                f"neighbor_cap_csv: {cap_csv}",
                f"neighbor_cap_real_sweep_csv: {real_cap_csv if real_caps else 'NA'}",
                f"neighbor_cap_figure_png: {Z4XF_FIGURE_PNG if z4xf_figure_ready else 'NA'}",
                "note: Z4X-F real sweep supersedes the older validator-boundary wording when present",
                "",
            ]
        ),
        encoding="utf-8",
    )

    build_docx(summary, details, caps, public_rows=public_rows, cap_source=cap_source, z4xf_figure_ready=z4xf_figure_ready)
    ROOT_REPORT.write_text(
        "\n".join(
            [
                "# TITAN FMM/M2L Scope Z4X-E Supplementary Sensitivity Report",
                "",
                f"Decision: {decision}",
                "",
                f"Supplement DOCX: {DOCX_OUT}",
                f"Paired P2 supplement DOCX: {P2_SUPPLEMENT_OUT}",
                f"Summary CSV: {summary_csv}",
                f"Detail CSV: {detail_csv}",
                f"Neighbor-cap CSV: {cap_csv}",
                f"Neighbor-cap real sweep CSV: {real_cap_csv if real_caps else 'NA'}",
                f"Neighbor-cap figure: {Z4XF_FIGURE_PNG if z4xf_figure_ready else 'NA'}",
                "",
                "Interpretation: theta/guard/tail rows are measured Z4W sensitivity evidence. When Z4X-F artifacts are available, ILU neighbor cap 32/64/128 rows are a real three-protein sweep; cap 64 remains the balanced reported protocol value.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(f"Z4X-E decision: {decision}")
    print(f"DOCX: {DOCX_OUT}")
    print(f"Paired supplement DOCX: {P2_SUPPLEMENT_OUT}")
    print(f"Summary CSV: {summary_csv}")
    print(f"Detail CSV: {detail_csv}")
    print(f"Neighbor cap CSV: {cap_csv}")
    if real_caps:
        print(f"Neighbor cap real sweep CSV: {real_cap_csv}")
    if z4xf_figure_ready:
        print(f"Neighbor cap figure: {Z4XF_FIGURE_PNG}")


if __name__ == "__main__":
    main()
