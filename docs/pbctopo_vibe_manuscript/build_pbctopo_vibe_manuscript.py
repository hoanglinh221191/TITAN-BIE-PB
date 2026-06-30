from __future__ import annotations

from pathlib import Path
from textwrap import dedent
import zipfile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "pbctopo_vibe_algorithm_manuscript.docx"
FIG_PATH = OUT_DIR / "pbctopo_vibe_algorithm_schematic.png"


INK = RGBColor(24, 36, 52)
BLUE = RGBColor(46, 116, 181)
TEAL = RGBColor(23, 113, 124)
MUTED = RGBColor(98, 109, 123)
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E8F4F5"
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"


def set_run_font(run, name: str = "Calibri", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 color: RGBColor | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tbl_grid = table._tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)


def table_text(cell, text: str, bold: bool = False, color: RGBColor = INK,
               size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True, color=MUTED)


def add_code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.4])
    cell = table.cell(0, 0)
    shade_cell(cell, "F7F8FA")
    set_cell_margins(cell, top=140, start=160, bottom=140, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line_idx, line in enumerate(code.strip("\n").splitlines()):
        if line_idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_run_font(r, "Consolas", 8.4, color=INK)


def add_callout(doc: Document, label: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.4])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + " ")
    set_run_font(r, bold=True, size=10.5, color=TEAL)
    r = p.add_run(body)
    set_run_font(r, size=10.5, color=INK)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else current + " " + word
        if draw.textbbox((0, 0), trial, font=font)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def rounded_box(draw, xy, fill, outline, label, body, title_font, body_font):
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.text((x0 + 22, y0 + 18), label, font=title_font, fill=outline)
    y = y0 + 56
    for line in wrap_text(draw, body, body_font, x1 - x0 - 44):
        draw.text((x0 + 22, y), line, font=body_font, fill=(35, 45, 58))
        y += 24


def arrow(draw, start, end, color=(46, 116, 181)):
    draw.line([start, end], fill=color, width=4)
    x0, y0 = start
    x1, y1 = end
    if abs(x1 - x0) >= abs(y1 - y0):
        sign = 1 if x1 >= x0 else -1
        pts = [(x1, y1), (x1 - sign * 16, y1 - 9), (x1 - sign * 16, y1 + 9)]
    else:
        sign = 1 if y1 >= y0 else -1
        pts = [(x1, y1), (x1 - 9, y1 - sign * 16), (x1 + 9, y1 - sign * 16)]
    draw.polygon(pts, fill=color)


def create_schematic() -> None:
    img = Image.new("RGB", (1800, 1040), "white")
    draw = ImageDraw.Draw(img)
    title_font = load_font(34, True)
    box_title = load_font(26, True)
    body_font = load_font(22)
    small_font = load_font(18)

    draw.text((70, 42), "VIBE / PBCTOPO correction pipeline", font=title_font,
              fill=(24, 36, 52))
    draw.text((70, 86),
              "Virtual Intact Block Extraction for topology-aware PBC reconstruction",
              font=small_font, fill=(98, 109, 123))

    boxes = [
        ((80, 160, 390, 335), (232, 238, 245), (46, 116, 181),
         "Inputs", "Periodic frames, two molecular systems, box matrix, topology, groups, optional warm state."),
        ((520, 160, 830, 335), (232, 244, 245), (23, 113, 124),
         "Topology context", "Heavy atoms, local edges, chain ranges, continuity pairs, anchor subset, owner contacts."),
        ((960, 160, 1270, 335), (247, 248, 250), (86, 98, 112),
         "Split graph", "Build local components using bond cutoff; choose exact, merge, or beam mode."),
        ((80, 535, 390, 735), (255, 248, 230), (122, 90, 0),
         "Candidate search", "Virtual cut planes plus component-shift graph search propose intact placements."),
        ((520, 535, 830, 735), (247, 248, 250), (86, 98, 112),
         "Score guards", "Reject clashes; preserve continuity; compare anchor graph, compactness, and shape metrics."),
        ((960, 535, 1270, 735), (232, 238, 245), (46, 116, 181),
         "Accept / refine", "Warm start, local refinement, split-owner guard, and best-safe candidate selection."),
        ((1395, 535, 1705, 735), (232, 244, 245), (23, 113, 124),
         "Apply frame", "Write corrected coordinates. Use pbci only as fallback when VIBE cannot certify a frame."),
    ]
    for xy, fill, outline, label, body in boxes:
        rounded_box(draw, xy, fill, outline, label, body, box_title, body_font)

    arrow(draw, (390, 248), (520, 248))
    arrow(draw, (830, 248), (960, 248))
    arrow(draw, (1115, 335), (235, 535), (122, 90, 0))
    arrow(draw, (390, 635), (520, 635), (122, 90, 0))
    arrow(draw, (830, 635), (960, 635))
    arrow(draw, (1270, 635), (1395, 635))

    draw.rounded_rectangle((84, 830, 1710, 950), radius=14, fill=(248, 250, 252),
                           outline=(210, 218, 226), width=2)
    legend = [
        (120, 360, "Cutoffs", "bond, clash, contact, continuity"),
        (520, 360, "Combinatorics", "exact component limit, beam width"),
        (920, 360, "Anchor backend", "auto, prim, spatial-all"),
        (1320, 330, "Baseline relation", "VIBE is independent; FixBox (-pbcn) and pbci are not extensions of VIBE"),
    ]
    for x, width, head, body in legend:
        draw.text((x, 850), head, font=load_font(20, True), fill=(46, 116, 181))
        y = 878
        for line in wrap_text(draw, body, load_font(18), width):
            draw.text((x, y), line, font=load_font(18), fill=(35, 45, 58))
            y += 22

    img.save(FIG_PATH, quality=95)


def add_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("VIBE: Virtual Intact Block Extraction for Topology-Aware PBC Correction")
    set_run_font(r, size=20, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("A manuscript-ready method note for TITAN pbctopo")
    set_run_font(r, size=11, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Draft status: algorithm description and implementation manuscript scaffold. "
                  "Benchmark claims are intentionally deferred until the external PDB/XTC test drive is mounted.")
    set_run_font(r, size=9.5, color=MUTED)


def add_parameter_table(doc: Document) -> None:
    doc.add_heading("Runtime Parameters Exposed Through CLI/Config", level=2)
    rows = [
        ("Parameter", "Default", "MARTINI default", "Algorithmic role"),
        ("bond cutoff", "2.15", "5.5", "Builds the local topology graph used to split atoms into connected components."),
        ("clash cutoff", "1.05", "3.5", "Rejects placements that create hard overlaps; also sets the legacy hard-clash guard."),
        ("contact cutoff", "4.5", "8.0", "Defines owner/contact relationships used by split guards and scoring context."),
        ("continuity cutoff", "20.0", "20.0", "Bounds continuity graph checks across frames and candidate placements."),
        ("exact component limit", "6", "6", "Controls when VIBE uses exact component-offset search instead of large-mode paths."),
        ("beam width", "12", "12", "Controls retained states in large-component beam search."),
        ("anchor backend", "auto", "auto", "Chooses anchor scoring backend: auto, prim, or spatial-all."),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    set_table_width(table, [1.55, 1.05, 1.35, 2.45])
    for ridx, row in enumerate(rows):
        for cidx, text in enumerate(row):
            cell = table.cell(ridx, cidx)
            table_text(cell, text, bold=(ridx == 0), size=9.0 if ridx else 9.3)
            shade_cell(cell, LIGHT_BLUE if ridx == 0 else WHITE)
    add_caption(doc, "Table 1. Public VIBE control surface after the pbctopo refactor.")


def add_pseudocode(doc: Document) -> None:
    doc.add_heading("Algorithm 1. VIBE / PBCTOPO", level=2)
    code = r"""
Input: frames F, receptor system R, ligand system L, periodic box B,
       topology T, runtime cutoffs C, exact_limit E, beam_width W
Output: corrected coordinates for each frame, plus diagnostics

Build static topology context from R, L, and T
Identify heavy atoms, continuity pairs, anchors, chain ranges
Initialize warm_state from the previous accepted frame, if available

for each frame f in F:
    atoms <- append receptor and ligand atoms for frame f
    local_edges <- topology graph using C.bond_cutoff
    components <- connected_components(local_edges)
    owner_contacts <- contacts(atoms, C.contact_cutoff)

    candidate_set <- empty
    if warm_state has a previous virtual-box candidate:
        score it first and keep it if safe

    projection <- heavy-atom projection plus local-edge subset
    seeds <- virtual cut-plane density seeds along x, y, z
    for each seed in seeds:
        coarse <- evaluate virtual-box placement on projection
        if coarse is promising:
            candidate_set.add(verify on full atoms)

    mode <- choose_component_mode(count(components), E)
    if mode == exact:
        graph_candidate <- exact component-offset graph search
    else if mode == merge:
        graph_candidate <- chain-aware merged-component graph search
    else:
        graph_candidate <- beam component search with width W
    candidate_set.add(graph_candidate)

    for each candidate in candidate_set:
        reject if hard clash under C.clash_cutoff
        reject if continuity graph fails under C.continuity_cutoff
        reject if anchor graph is disconnected or worse than incumbent
        score compactness, split-owner safety, shape, and shift norm

    best <- best safe candidate
    if best is valid:
        apply best to frame f
        update warm_state
    else:
        apply pbci fallback for frame f and record fallback diagnostic
"""
    add_code_block(doc, code)
    p = doc.add_paragraph()
    r = p.add_run("Note. ")
    set_run_font(r, bold=True, color=TEAL)
    r = p.add_run("The pbci step is a fallback/baseline path, not part of the VIBE search model. "
                  "FixBox (-pbcn) is likewise a separate TITAN tool option.")
    set_run_font(r, color=INK)


def add_manuscript_text(doc: Document) -> None:
    add_title(doc)

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Periodic boundary condition artifacts can fragment molecular assemblies, "
        "obscure chain continuity, and distort downstream interaction analyses. "
        "We describe VIBE (Virtual Intact Block Extraction), implemented in TITAN "
        "as pbctopo, a topology-aware reconstruction strategy that searches for "
        "intact molecular placements using virtual cut-plane candidates, component "
        "offset graphs, continuity constraints, and anchor-based scoring. The method "
        "is designed for multi-chain protein systems and coarse-grained MARTINI "
        "trajectories while keeping conservative fallback behavior when a frame "
        "cannot be certified. This manuscript scaffold documents the algorithm, "
        "public configuration surface, implementation boundaries, and validation "
        "plan without making benchmark claims that require the currently unmounted "
        "external PDB/XTC test drive."
    )

    add_callout(
        doc,
        "Contribution summary:",
        "VIBE is presented as an independent pbctopo algorithm in TITAN. "
        "FixBox (-pbcn) and pbci are described only as separate TITAN baseline or "
        "fallback paths, not as VIBE extensions.",
        LIGHT_TEAL,
    )

    doc.add_heading("1. Motivation", level=1)
    doc.add_paragraph(
        "Many MD trajectories store atoms inside a primary periodic image even when "
        "the chemically meaningful assembly spans an image boundary. Simple minimum "
        "image wrapping can restore distances for pair calculations, but it may not "
        "recover an intact block suitable for output PDBs, contact maps, radius of "
        "gyration calculations, or multi-chain topology analyses. VIBE targets this "
        "gap: it searches for an atom placement that is simultaneously compact, "
        "topology-consistent, continuity-preserving, and guarded against hard clashes."
    )
    doc.add_paragraph(
        "The method is intentionally conservative. It does not declare a corrected "
        "frame merely because a virtual split looks plausible. Candidate placements "
        "must pass local topology checks, owner-contact split guards, continuity "
        "constraints, anchor graph scoring, and clash rejection. If this certification "
        "fails, TITAN records the failure and can fall back to the existing pbci path."
    )

    doc.add_heading("2. Method Overview", level=1)
    doc.add_paragraph(
        "VIBE operates per frame but carries warm-state information forward from "
        "previous accepted frames. For each frame, TITAN extracts receptor and ligand "
        "atoms, builds a local topology graph, identifies heavy-atom projection "
        "subsets, derives continuity and anchor pairs, and then searches candidate "
        "placements through two complementary mechanisms."
    )
    doc.add_paragraph(
        "The first mechanism is virtual-box search. Candidate cut planes are generated "
        "from density and gap structure along the periodic axes, evaluated cheaply on "
        "a projected subset, and then verified against the full atom set. The second "
        "mechanism is component-offset graph search. Local topology splits atoms into "
        "connected components; small component counts use exact search, intermediate "
        "counts can merge chain-local raw components before graph solving, and large "
        "counts use a bounded beam search. The selected mode is controlled by the "
        "exact component limit and beam width exposed through the pbctopo CLI/config."
    )

    doc.add_picture(str(FIG_PATH), width=Inches(6.4))
    add_caption(doc, "Figure 1. Schematic of the VIBE/pbctopo frame-correction pipeline.")

    add_pseudocode(doc)

    doc.add_heading("3. Scoring and Rejection Criteria", level=1)
    doc.add_paragraph(
        "Candidate scoring is multi-stage. A placement first has to avoid ultra-clash "
        "conditions under the clash cutoff. It then has to preserve continuity graph "
        "metrics under the continuity cutoff and keep anchor relationships connected "
        "under the selected anchor backend. Surviving candidates are compared using "
        "shift norm, compactness, bounding-box and gyration-derived shape statistics, "
        "owner-pair split safety, and whether the candidate improves or remains close "
        "to the current incumbent. This layered design is deliberate: the algorithm "
        "prefers a certified safe placement over a geometrically attractive but "
        "topologically suspicious placement."
    )

    add_parameter_table(doc)

    doc.add_heading("4. Implementation Surface in TITAN", level=1)
    doc.add_paragraph(
        "The public command is -pbctopo. The refactor introduces a pbctopo module "
        "boundary under src/pbctopo for runtime configuration and logging while "
        "keeping Multi_MOL::solve_pbc_pbctopo as the stable public adapter. Runtime "
        "cutoffs are resolved after force-field defaults are selected, so MARTINI "
        "defaults remain intact unless a user explicitly overrides them. The same "
        "configuration fields are documented in CLI help and exported through the "
        "TITAN schema surface."
    )
    add_callout(
        doc,
        "Current refactor boundary:",
        "Configuration and logging are already isolated in src/pbctopo. The next "
        "low-risk extractions should be geometry helpers, topology-context builders, "
        "scoring, component search, and virtual-box evaluation.",
        LIGHT_BLUE,
    )

    doc.add_heading("5. Validation Plan", level=1)
    validation = [
        "Parser and schema contract: invalid cutoffs, invalid beam/exact limits, and invalid anchor backend fail closed.",
        "Smoke behavior: existing PBC/nocal smoke remains unchanged when VIBE options are absent.",
        "Frame-level diagnostics: pbctopo_debug.log should record resolved cutoffs, anchor backend, exact limit, beam width, component modes, and fallback counts.",
        "Regression windows: rerun dimer and tetramer pbctopo windows once the external D:\\Data_MD drive is mounted.",
        "Publication metrics: report frame success rate, pbci fallback count, component graph usage, runtime, and geometry sanity checks before making benchmark claims.",
    ]
    for item in validation:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, color=INK)

    doc.add_heading("6. Limitations and Manuscript Notes", level=1)
    doc.add_paragraph(
        "This draft should be treated as a methods manuscript scaffold rather than a "
        "final results paper. It describes the algorithm and implementation state, "
        "but it deliberately avoids new quantitative claims because the PDB/XTC "
        "validation drive is not currently mounted. Once the drive is available, the "
        "existing dimer and tetramer windows can supply method validation tables and "
        "representative corrected structures."
    )
    doc.add_paragraph(
        "Terminology should stay strict in the final manuscript: VIBE is the pbctopo "
        "algorithm, FixBox is the separate -pbcn option, and pbci is an in-house TITAN "
        "PBC path used here only as a baseline or fallback reference."
    )

    doc.add_heading("Acknowledgments and Disclosure Draft", level=1)
    doc.add_paragraph(
        "This manuscript scaffold was generated with AI-assisted drafting from the "
        "local TITAN source tree. All algorithmic claims should be reviewed against "
        "the source code and rerun logs before submission. No external benchmark data "
        "were analyzed in this scaffold."
    )


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "VIBE: Virtual Intact Block Extraction for Topology-Aware PBC Correction"
    props.subject = "TITAN pbctopo algorithm manuscript scaffold"
    props.author = "TITAN development team"
    props.keywords = "TITAN, pbctopo, VIBE, periodic boundary correction, molecular dynamics"


def structural_audit(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path) as zf:
        names = set(zf.namelist())
        required = {
            "word/document.xml",
            "word/media/image1.png",
            "docProps/core.xml",
        }
        missing = required - names
        if missing:
            raise RuntimeError(f"Missing DOCX parts: {sorted(missing)}")
        xml = zf.read("word/document.xml").decode("utf-8")
        for needle in [
            "Virtual Intact Block Extraction",
            "Algorithm 1. VIBE",
            "Runtime Parameters Exposed",
            "FixBox (-pbcn)",
            "pbci fallback",
        ]:
            if needle not in xml:
                raise RuntimeError(f"Missing expected manuscript text: {needle}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_schematic()
    doc = Document()
    add_styles(doc)
    set_core_properties(doc)
    add_manuscript_text(doc)
    doc.save(DOCX_PATH)
    structural_audit(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")
    print(f"Wrote {FIG_PATH}")


if __name__ == "__main__":
    main()
